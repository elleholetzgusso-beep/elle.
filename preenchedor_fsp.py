#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
import openpyxl

# ======================= CONFIG (ajusta aqui) =======================
FSP_PATH        = "FSP.xlsx"
SHEET_NAME      = "Control doc. Aportados"
DOCS_FOLDER     = "documentos"

HEADER_ROW      = 1
COL_CODIGO      = "A"   # Id. ([DE-1], [DE-2]...)
COL_DESCRIPCION = "C"   # Descripcion
COL_REF         = "B"   # Ref. Documento  <-- o script preenche aqui

SKIP_PREFIXES   = []
SKIP_CODE_LIST  = []    # ex.: ["[DE-34]", "[DE-35]"] p/ pular planos

OVERWRITE_FILLED = False
MATCH_THRESHOLD  = 0.55
VERSION_REGEX    = r'[_\-\s]?[vV](\d{1,3})'
# ====================================================================


def normalizar(txt):
    if txt is None:
        return ""
    txt = str(txt)
    txt = unicodedata.normalize("NFKD", txt).encode("ascii", "ignore").decode()
    txt = re.sub(r"[^a-z0-9]+", " ", txt.lower())
    return re.sub(r"\s+", " ", txt).strip()

def similaridade(a, b):
    return SequenceMatcher(None, a, b).ratio()

def indexar_pasta(folder):
    itens = []
    for f in sorted(folder.iterdir()):
        if not f.is_file() or f.name.startswith("~$"):
            continue
        nome = f.stem
        m = re.search(VERSION_REGEX, nome)
        versao = f"v{int(m.group(1)):02d}" if m else None
        desc = re.sub(VERSION_REGEX, "", nome)
        itens.append((normalizar(desc), versao, f.name))
    return itens

def melhor_match(descricao_fsp, itens):
    alvo = normalizar(descricao_fsp)
    if not alvo:
        return None, 0.0
    melhor, score = None, 0.0
    for desc_norm, versao, nome in itens:
        s = similaridade(alvo, desc_norm)
        if s > score:
            melhor, score = (versao, nome), s
    return melhor, score

def col(letter):
    return openpyxl.utils.column_index_from_string(letter)

def main():
    fsp = Path(FSP_PATH)
    pasta = Path(DOCS_FOLDER)
    if not fsp.exists():
        raise SystemExit(f"ERRO: Nao encontrei o FSP: {fsp}")
    if not pasta.exists():
        raise SystemExit(f"ERRO: Nao encontrei a pasta de documentos: {pasta}")

    itens = indexar_pasta(pasta)
    print(f"{len(itens)} ficheiros lidos da pasta '{pasta}'.\n")

    wb = openpyxl.load_workbook(fsp)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"ERRO: Aba '{SHEET_NAME}' nao existe. Abas: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    c_cod, c_desc, c_ref = col(COL_CODIGO), col(COL_DESCRIPCION), col(COL_REF)
    preenchidas, puladas, sem_match = [], [], []

    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        codigo = ws.cell(row, c_cod).value
        desc   = ws.cell(row, c_desc).value
        if not desc and not codigo:
            continue

        cod_str = str(codigo).strip() if codigo else ""
        if cod_str in SKIP_CODE_LIST or any(cod_str.startswith(p) for p in SKIP_PREFIXES):
            puladas.append(cod_str or "(s/codigo)")
            continue

        atual = ws.cell(row, c_ref).value
        if atual and not OVERWRITE_FILLED:
            continue

        match, score = melhor_match(desc or "", itens)
        if match and score >= MATCH_THRESHOLD:
            versao, nome = match
            ref = f"{str(desc).strip()}_{versao}" if versao else str(desc).strip()
            ws.cell(row, c_ref).value = ref
            preenchidas.append((cod_str, ref, nome, round(score, 2)))
        else:
            sem_match.append((cod_str, desc, round(score, 2)))

    saida = fsp.with_name(fsp.stem + "_RELLENADO.xlsx")
    wb.save(saida)

    print(f"PREENCHIDAS ({len(preenchidas)}):")
    for cod, ref, nome, sc in preenchidas:
        print(f"   {cod:<10} -> {ref:<50} [arquivo: {nome}]  (sim={sc})")
    print(f"\nPULADAS ({len(puladas)}): {', '.join(puladas)}")
    print(f"\nSEM CORRESPONDENCIA - revisar a mao ({len(sem_match)}):")
    for cod, desc, sc in sem_match:
        print(f"   {cod:<10} '{desc}'  (melhor sim={sc})")
    print(f"\nGuardado em: {saida}")

if __name__ == "__main__":
    main()
