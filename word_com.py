"""Automatización de Microsoft Word mediante COM (pywin32).

Módulo de apoyo para la evaluación independiente de seguridad (AsBo/ISA) de
proyectos ferroviarios: comparación de versiones sucesivas de documentación
técnica, extracción de comentarios y de control de cambios, actualización de
campos e índices, y exportación a PDF con paginación real.

Requisitos
----------
* Windows con Microsoft Word instalado (este módulo pilota el Word real).
* pywin32:   pip install pywin32
* openpyxl:  pip install openpyxl      (opcional; sin él se exporta a CSV)

Principios de diseño
--------------------
1. El documento original NUNCA se modifica. Toda operación que escribe trabaja
   sobre una copia local en una carpeta temporal.
2. Word se abre siempre invisible y sin diálogos, en una instancia DEDICADA
   (DispatchEx), para no interferir con el Word que la usuaria tenga abierto.
3. Todo se cierra en bloques `finally`: ningún proceso WINWORD.EXE debe quedar
   colgado, ni siquiera si el script revienta.
4. Las rutas se convierten SIEMPRE a absolutas: Word resuelve las rutas
   relativas contra su propio directorio de trabajo, no contra el de Python.
"""

from __future__ import annotations

import csv
import logging
import re
import shutil
import subprocess
import tempfile
from contextlib import contextmanager
from dataclasses import asdict, dataclass, fields as campos_dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator, Sequence

try:  # pragma: no cover - depende del sistema operativo
    import pythoncom
    import win32com.client
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "Este módulo requiere pywin32 y Microsoft Word sobre Windows. "
        "Instálelo con: pip install pywin32"
    ) from exc


__all__ = [
    "SesionWord",
    "Comentario",
    "Revision",
    "ResultadoLote",
    "ErrorDocumento",
    "comparar_documentos",
    "extraer_comentarios",
    "extraer_revisiones",
    "actualizar_campos",
    "exportar_a_pdf",
    "comparar_lote",
    "extraer_comentarios_lote",
    "emparejar_versiones",
    "exportar_tabla",
    "configurar_registro",
    "diagnostico",
]


# ---------------------------------------------------------------------------
# Constantes de Word
# ---------------------------------------------------------------------------
# Se conservan los nombres originales de VBA (wdXxx) para poder buscarlos tal
# cual en la documentación de Microsoft Learn ("WdSaveFormat enumeration", etc.)
# o en el Examinador de objetos del editor VBA de Word (Alt+F11, luego F2).


class Wd:
    """Constantes numéricas del modelo de objetos de Word."""

    # WdSaveFormat
    wdFormatDocumentDefault = 16   # .docx
    wdFormatXMLDocument = 12       # .docx (equivalente)
    wdFormatPDF = 17               # .pdf

    # WdSaveOptions
    wdDoNotSaveChanges = 0
    wdSaveChanges = -1

    # WdCompareTarget
    wdCompareTargetSelected = 0
    wdCompareTargetCurrent = 1
    wdCompareTargetNew = 2         # el resultado va a un documento nuevo

    # WdGranularity
    wdGranularityCharLevel = 0
    wdGranularityWordLevel = 1

    # WdInformation
    wdActiveEndAdjustedPageNumber = 1
    wdActiveEndPageNumber = 3
    wdNumberOfPagesInDocument = 4

    # WdExportFormat / ExportAsFixedFormat
    wdExportFormatPDF = 17
    wdExportOptimizeForPrint = 0
    wdExportOptimizeForOnScreen = 1
    wdExportAllDocument = 0
    wdExportDocumentContent = 0    # PDF sin marcas de revisión
    wdExportDocumentWithMarkup = 7  # PDF con el control de cambios visible
    wdExportCreateNoBookmarks = 0
    wdExportCreateHeadingBookmarks = 1

    # WdProtectionType
    wdNoProtection = -1

    # MsoAutomationSecurity
    msoAutomationSecurityForceDisable = 3  # abre sin ejecutar macros

    # WdRevisionType
    wdRevisionInsert = 1
    wdRevisionDelete = 2
    wdRevisionProperty = 3
    wdRevisionParagraphNumber = 4
    wdRevisionDisplayField = 5
    wdRevisionReconcile = 6
    wdRevisionConflict = 7
    wdRevisionStyle = 8
    wdRevisionReplace = 9
    wdRevisionParagraphProperty = 10
    wdRevisionTableProperty = 11
    wdRevisionSectionProperty = 12
    wdRevisionStyleDefinition = 13
    wdRevisionMovedFrom = 14
    wdRevisionMovedTo = 15
    wdRevisionCellInsertion = 16
    wdRevisionCellDeletion = 17
    wdRevisionCellMerge = 18


#: Descripción en castellano de cada tipo de revisión (WdRevisionType).
TIPOS_DE_REVISION: dict[int, str] = {
    Wd.wdRevisionInsert: "Inserción",
    Wd.wdRevisionDelete: "Eliminación",
    Wd.wdRevisionProperty: "Cambio de formato",
    Wd.wdRevisionParagraphNumber: "Numeración de párrafo",
    Wd.wdRevisionDisplayField: "Campo mostrado",
    Wd.wdRevisionReconcile: "Reconciliación",
    Wd.wdRevisionConflict: "Conflicto",
    Wd.wdRevisionStyle: "Cambio de estilo",
    Wd.wdRevisionReplace: "Sustitución",
    Wd.wdRevisionParagraphProperty: "Formato de párrafo",
    Wd.wdRevisionTableProperty: "Formato de tabla",
    Wd.wdRevisionSectionProperty: "Formato de sección",
    Wd.wdRevisionStyleDefinition: "Definición de estilo",
    Wd.wdRevisionMovedFrom: "Texto movido (origen)",
    Wd.wdRevisionMovedTo: "Texto movido (destino)",
    Wd.wdRevisionCellInsertion: "Inserción de celda",
    Wd.wdRevisionCellDeletion: "Eliminación de celda",
    Wd.wdRevisionCellMerge: "Combinación de celdas",
}

#: Patrón de versión al final del nombre de fichero: "..._v03", "...-V4".
PATRON_VERSION = re.compile(r"^(?P<base>.+?)[ _\-]*[vV](?P<num>\d{1,3})$")

registro = logging.getLogger("word_com")


# ---------------------------------------------------------------------------
# Excepciones y estructuras de datos
# ---------------------------------------------------------------------------


class ErrorDocumento(Exception):
    """Fallo al procesar un documento concreto.

    Se usa para que un documento protegido, bloqueado por otro proceso o
    corrupto no aborte el procesamiento del resto del lote.
    """

    def __init__(self, ruta: Path | str, motivo: str) -> None:
        self.ruta = Path(ruta)
        self.motivo = motivo
        super().__init__(f"{self.ruta.name}: {motivo}")


@dataclass(frozen=True)
class Comentario:
    """Un comentario extraído de un documento de Word."""

    documento: str
    indice: int
    autor: str
    iniciales: str
    fecha: str
    pagina: int
    apartado: str
    texto_comentario: str
    texto_comentado: str
    resuelto: str


@dataclass(frozen=True)
class Revision:
    """Un cambio registrado (control de cambios) de un documento de Word."""

    documento: str
    indice: int
    autor: str
    fecha: str
    tipo: int
    descripcion_tipo: str
    pagina: int
    apartado: str
    estilo: str
    texto: str


@dataclass
class ResultadoLote:
    """Resumen de una ejecución en lote."""

    correctos: list[str]
    fallidos: list[tuple[str, str]]
    salidas: list[Path]

    @property
    def total(self) -> int:
        return len(self.correctos) + len(self.fallidos)

    def resumen(self) -> str:
        """Devuelve un resumen legible en castellano."""
        lineas = [
            f"Procesados: {self.total} | Correctos: {len(self.correctos)} "
            f"| Con incidencias: {len(self.fallidos)}"
        ]
        for nombre, motivo in self.fallidos:
            lineas.append(f"  - INCIDENCIA en {nombre}: {motivo}")
        return "\n".join(lineas)


# ---------------------------------------------------------------------------
# Utilidades internas
# ---------------------------------------------------------------------------


def _ruta_absoluta(ruta: Path | str) -> Path:
    """Normaliza una ruta a absoluta.

    Word resuelve las rutas relativas contra SU directorio de trabajo (que es
    imprevisible y cambia al abrir o guardar ficheros), no contra el de Python.
    Una ruta relativa provoca o bien un error de "no se encuentra el archivo",
    o bien —peor— la apertura silenciosa de un fichero distinto.
    """
    return Path(ruta).expanduser().resolve()


def _texto_limpio(texto: str | None, maximo: int | None = None) -> str:
    """Limpia los caracteres de control propios de Word.

    Word devuelve marcas de fin de párrafo (\\r), de celda (\\x07) y saltos de
    línea manuales (\\x0b) dentro del texto de un Range.
    """
    if not texto:
        return ""
    limpio = texto.replace("\r", " ").replace("\x07", " ").replace("\x0b", " ")
    limpio = limpio.replace("\a", " ").replace("\n", " ")
    limpio = re.sub(r"\s+", " ", limpio).strip()
    if maximo is not None and len(limpio) > maximo:
        limpio = limpio[:maximo].rstrip() + "…"
    return limpio


def _texto_fecha(valor: Any) -> str:
    """Convierte una fecha COM (pywintypes) a texto en formato español."""
    if valor is None:
        return ""
    try:
        return datetime(
            valor.year, valor.month, valor.day,
            valor.hour, valor.minute, valor.second,
        ).strftime("%d/%m/%Y %H:%M")
    except (AttributeError, ValueError, TypeError):
        return str(valor)


def _atributo_seguro(objeto: Any, nombre: str, defecto: Any = "") -> Any:
    """Lee una propiedad COM que puede no existir en esta versión de Word."""
    try:
        valor = getattr(objeto, nombre)
    except Exception:  # noqa: BLE001 - COM lanza excepciones muy variadas
        return defecto
    return defecto if valor is None else valor


def _pagina_de(rango: Any) -> int:
    """Devuelve el número de página real de un Range.

    La paginación no existe en el fichero .docx: la calcula el motor de
    maquetación de Word. Por eso esta información sólo se puede obtener con el
    documento abierto en Word, y no con python-docx.
    """
    try:
        return int(rango.Information(Wd.wdActiveEndPageNumber))
    except Exception:  # noqa: BLE001
        return 0


def _apartado_de(rango: Any) -> str:
    """Devuelve la numeración automática del párrafo (p. ej. "3.2.1")."""
    try:
        return _texto_limpio(rango.Paragraphs(1).Range.ListFormat.ListString)
    except Exception:  # noqa: BLE001
        return ""


def _estilo_de(rango: Any) -> str:
    """Devuelve el nombre del estilo del párrafo que contiene el Range."""
    try:
        return str(rango.Paragraphs(1).Range.Style.NameLocal)
    except Exception:  # noqa: BLE001
        return ""


def _es_documento_word(ruta: Path) -> bool:
    """Indica si el fichero es un documento de Word procesable.

    Descarta los ficheros temporales de bloqueo que Word crea junto al
    original ("~$informe.docx") y que no son documentos reales.
    """
    if ruta.name.startswith("~$"):
        return False
    return ruta.suffix.lower() in {".doc", ".docx", ".docm", ".rtf"}


def configurar_registro(nivel: int = logging.INFO) -> None:
    """Configura la salida por consola del registro de actividad."""
    logging.basicConfig(
        level=nivel,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        datefmt="%H:%M:%S",
    )


# ---------------------------------------------------------------------------
# Sesión de Word
# ---------------------------------------------------------------------------


class SesionWord:
    """Gestor de contexto para una instancia dedicada de Microsoft Word.

    Uso::

        with SesionWord() as sesion:
            with sesion.abrir(r"C:\\ruta\\informe.docx") as doc:
                print(doc.Paragraphs.Count)

    Al salir del bloque `with` se cierran todos los documentos y se termina el
    proceso WINWORD.EXE, ocurra lo que ocurra dentro del bloque.

    Se emplea DispatchEx y no Dispatch: Dispatch se engancha a la instancia de
    Word que ya esté abierta, con lo que el `Quit()` final cerraría el Word en
    el que la usuaria esté trabajando. DispatchEx arranca un proceso nuevo y
    exclusivo de este script.
    """

    def __init__(
        self,
        visible: bool = False,
        mostrar_alertas: bool = False,
        desactivar_macros: bool = True,
    ) -> None:
        self.visible = visible
        self.mostrar_alertas = mostrar_alertas
        self.desactivar_macros = desactivar_macros
        self._app: Any | None = None
        self._pid: int | None = None
        self._com_iniciado = False

    # -- ciclo de vida ----------------------------------------------------

    def __enter__(self) -> "SesionWord":
        self.iniciar()
        return self

    def __exit__(self, *_excepcion: object) -> None:
        self.cerrar()

    def iniciar(self) -> None:
        """Arranca una instancia dedicada de Word."""
        pythoncom.CoInitialize()
        self._com_iniciado = True
        registro.debug("Arrancando instancia dedicada de Microsoft Word…")
        self._app = win32com.client.DispatchEx("Word.Application")
        self._app.Visible = self.visible
        # Sin esto, cualquier diálogo modal (archivo bloqueado, conversión de
        # formato, guardar cambios) dejaría el script colgado para siempre:
        # la ventana no es visible y no hay nadie que pueda pulsar "Aceptar".
        self._app.DisplayAlerts = False
        try:
            self._app.ScreenUpdating = False
        except Exception:  # noqa: BLE001
            pass
        if self.desactivar_macros:
            try:
                self._app.AutomationSecurity = Wd.msoAutomationSecurityForceDisable
            except Exception:  # noqa: BLE001
                registro.debug("No se ha podido fijar AutomationSecurity.")
        self._pid = self._obtener_pid()
        registro.info("Word iniciado (versión %s, PID %s).",
                      _atributo_seguro(self._app, "Version", "desconocida"),
                      self._pid if self._pid else "desconocido")

    def cerrar(self) -> None:
        """Cierra los documentos y termina el proceso de Word.

        Nunca lanza excepción: es el método del bloque `finally`, y un fallo
        aquí enmascararía el error real que provocó la salida.
        """
        if self._app is not None:
            self._cerrar_documentos_abiertos()
            self._salir_de_word()
        self._app = None
        if self._com_iniciado:
            try:
                pythoncom.CoUninitialize()
            except Exception:  # noqa: BLE001
                pass
            self._com_iniciado = False

    def _cerrar_documentos_abiertos(self) -> None:
        """Cierra sin guardar todo lo que haya quedado abierto."""
        try:
            while self._app.Documents.Count > 0:
                self._app.Documents.Item(1).Close(Wd.wdDoNotSaveChanges)
        except Exception as exc:  # noqa: BLE001
            registro.warning("No se han podido cerrar todos los documentos: %s", exc)

    def _salir_de_word(self) -> None:
        """Termina Word; si la salida ordenada falla, mata el proceso."""
        try:
            self._app.Quit(Wd.wdDoNotSaveChanges)
            registro.info("Word cerrado correctamente.")
            return
        except Exception as exc:  # noqa: BLE001
            registro.warning("Word no ha respondido al cierre ordenado: %s", exc)
        if self._pid:
            self.matar_proceso()

    def matar_proceso(self) -> None:
        """Termina por la fuerza el proceso de Word de ESTA sesión.

        Red de seguridad de último recurso. Sólo afecta al PID arrancado por
        esta sesión: nunca a otras instancias de Word abiertas por la usuaria.
        """
        if not self._pid:
            registro.warning("No se conoce el PID de Word; no se puede forzar el cierre.")
            return
        registro.warning("Forzando el cierre del proceso WINWORD.EXE (PID %s).", self._pid)
        try:
            subprocess.run(
                ["taskkill", "/F", "/PID", str(self._pid)],
                capture_output=True,
                check=False,
            )
        except Exception as exc:  # noqa: BLE001
            registro.error("No se ha podido forzar el cierre del proceso: %s", exc)

    def _obtener_pid(self) -> int | None:
        """Obtiene el PID del proceso de Word a partir del identificador de ventana."""
        try:
            import win32process

            _, pid = win32process.GetWindowThreadProcessId(self._app.Hwnd)
            return int(pid)
        except Exception:  # noqa: BLE001
            return None

    # -- acceso -----------------------------------------------------------

    @property
    def app(self) -> Any:
        """Objeto Application de Word (la raíz del modelo de objetos)."""
        if self._app is None:
            raise RuntimeError("La sesión de Word no está iniciada.")
        return self._app

    @contextmanager
    def abrir(
        self,
        ruta: Path | str,
        solo_lectura: bool = True,
        contrasena: str = "",
    ) -> Iterator[Any]:
        """Abre un documento y garantiza su cierre.

        Args:
            ruta: Ruta del documento. Se convierte a absoluta.
            solo_lectura: Abrir en modo de sólo lectura (protección adicional
                del original, además de trabajar siempre sobre copias).
            contrasena: Contraseña de apertura, si el documento la tiene.

        Yields:
            El objeto Document de Word.

        Raises:
            ErrorDocumento: Si el documento no se puede abrir (inexistente,
                bloqueado por otro proceso, dañado o protegido).
        """
        ruta_abs = _ruta_absoluta(ruta)
        if not ruta_abs.is_file():
            raise ErrorDocumento(ruta_abs, "el fichero no existe o no es accesible")

        documento = None
        try:
            try:
                documento = self.app.Documents.Open(
                    FileName=str(ruta_abs),
                    ConfirmConversions=False,
                    ReadOnly=solo_lectura,
                    AddToRecentFiles=False,
                    PasswordDocument=contrasena,
                    Visible=False,
                )
            except Exception as exc:  # noqa: BLE001
                raise ErrorDocumento(ruta_abs, f"no se ha podido abrir ({exc})") from exc

            self._avisar_si_protegido(documento, ruta_abs)
            yield documento
        finally:
            # Se ejecuta siempre: también si el cuerpo del `with` lanza una
            # excepción. Es lo que impide que un documento defectuoso deje el
            # fichero bloqueado y aborte el resto del lote.
            if documento is not None:
                try:
                    documento.Close(Wd.wdDoNotSaveChanges)
                except Exception as exc:  # noqa: BLE001
                    registro.warning("Cierre incompleto de %s: %s", ruta_abs.name, exc)

    def _avisar_si_protegido(self, documento: Any, ruta: Path) -> None:
        """Registra un aviso si el documento tiene protección activa."""
        proteccion = _atributo_seguro(documento, "ProtectionType", Wd.wdNoProtection)
        if proteccion != Wd.wdNoProtection:
            registro.warning(
                "El documento %s está protegido (tipo %s); algunas operaciones "
                "pueden fallar.", ruta.name, proteccion,
            )


@contextmanager
def _carpeta_temporal() -> Iterator[Path]:
    """Crea una carpeta temporal local y la elimina al terminar."""
    carpeta = Path(tempfile.mkdtemp(prefix="word_com_"))
    try:
        yield carpeta
    finally:
        shutil.rmtree(carpeta, ignore_errors=True)


def _copia_de_trabajo(origen: Path, destino_carpeta: Path) -> Path:
    """Copia el documento a una carpeta local de trabajo.

    Cumple dos objetivos a la vez: garantiza que el original jamás se modifica
    y evita la Vista Protegida que Word aplica a los ficheros procedentes de
    unidades de red, que bloquea la edición y la comparación.
    """
    destino = destino_carpeta / origen.name
    shutil.copy2(origen, destino)
    return destino


# ---------------------------------------------------------------------------
# 1. Comparación de versiones
# ---------------------------------------------------------------------------


def comparar_documentos(
    sesion: SesionWord,
    ruta_original: Path | str,
    ruta_revisado: Path | str,
    ruta_salida: Path | str,
    autor_revision: str = "Comparación automática",
    nivel_palabra: bool = True,
    comparar_formato: bool = True,
) -> Path:
    """Compara dos versiones de un documento y guarda el resultado marcado.

    Emplea Application.CompareDocuments, el mismo algoritmo que "Revisar >
    Comparar" en la interfaz de Word. El resultado es un documento NUEVO en el
    que las diferencias aparecen como cambios registrados (control de cambios).
    Ninguno de los dos documentos de entrada se modifica.

    Args:
        sesion: Sesión de Word activa.
        ruta_original: Versión anterior (p. ej. ..._v03.docx).
        ruta_revisado: Versión nueva (p. ej. ..._v04.docx).
        ruta_salida: Ruta del .docx comparado que se va a generar.
        autor_revision: Nombre que figurará como autor de las marcas.
        nivel_palabra: True compara palabra a palabra (comportamiento estándar
            de Word); False compara carácter a carácter, más ruidoso.
        comparar_formato: Incluir también los cambios de formato.

    Returns:
        La ruta del documento comparado.

    Raises:
        ErrorDocumento: Si alguno de los documentos no se puede procesar.
    """
    origen_a = _ruta_absoluta(ruta_original)
    origen_b = _ruta_absoluta(ruta_revisado)
    salida = _ruta_absoluta(ruta_salida)
    salida.parent.mkdir(parents=True, exist_ok=True)

    registro.info("Comparando: %s  ->  %s", origen_a.name, origen_b.name)

    with _carpeta_temporal() as temporal:
        # Cada versión va a su propia subcarpeta: los dos ficheros pueden
        # llamarse igual si proceden de carpetas distintas.
        copia_a = _copia_de_trabajo_en(origen_a, temporal / "anterior")
        copia_b = _copia_de_trabajo_en(origen_b, temporal / "nueva")
        with sesion.abrir(copia_a, solo_lectura=False) as doc_a:
            with sesion.abrir(copia_b, solo_lectura=False) as doc_b:
                _avisar_de_cambios_pendientes(doc_a, origen_a)
                _avisar_de_cambios_pendientes(doc_b, origen_b)
                documento_comparado = _ejecutar_comparacion(
                    sesion, doc_a, doc_b, autor_revision,
                    nivel_palabra, comparar_formato,
                )
                _guardar_y_cerrar(documento_comparado, salida)

    registro.info("Comparativa generada: %s", salida.name)
    return salida


def _copia_de_trabajo_en(origen: Path, subcarpeta: Path) -> Path:
    """Variante de _copia_de_trabajo que crea la subcarpeta destino."""
    subcarpeta.mkdir(parents=True, exist_ok=True)
    return _copia_de_trabajo(origen, subcarpeta)


def _avisar_de_cambios_pendientes(documento: Any, ruta: Path) -> None:
    """Avisa si el documento llega con control de cambios sin aceptar.

    Word considera aceptadas las marcas pendientes al comparar. Se deja
    constancia en el registro para que quede documentado en el expediente.
    """
    try:
        pendientes = int(documento.Revisions.Count)
    except Exception:  # noqa: BLE001
        return
    if pendientes:
        registro.warning(
            "%s contiene %d cambios registrados pendientes. Word los tratará "
            "como aceptados durante la comparación.", ruta.name, pendientes,
        )


def _ejecutar_comparacion(
    sesion: SesionWord,
    doc_a: Any,
    doc_b: Any,
    autor_revision: str,
    nivel_palabra: bool,
    comparar_formato: bool,
) -> Any:
    """Invoca Application.CompareDocuments y devuelve el documento resultante."""
    granularidad = (
        Wd.wdGranularityWordLevel if nivel_palabra else Wd.wdGranularityCharLevel
    )
    try:
        return sesion.app.CompareDocuments(
            OriginalDocument=doc_a,
            RevisedDocument=doc_b,
            Destination=Wd.wdCompareTargetNew,
            Granularity=granularidad,
            CompareFormatting=comparar_formato,
            CompareCaseChanges=True,
            CompareWhitespace=True,
            CompareTables=True,
            CompareHeaders=True,
            CompareFootnotes=True,
            CompareTextboxes=True,
            CompareFields=True,
            CompareComments=True,
            CompareMoves=True,
            RevisedAuthor=autor_revision,
            IgnoreAllComparisonWarnings=True,
        )
    except Exception as exc:  # noqa: BLE001
        raise ErrorDocumento(
            doc_b.FullName, f"Word no ha podido comparar los documentos ({exc})"
        ) from exc


def _guardar_y_cerrar(documento: Any, salida: Path) -> None:
    """Guarda un documento como .docx en la ruta indicada y lo cierra."""
    try:
        _guardar_como(documento, salida, Wd.wdFormatDocumentDefault)
    finally:
        try:
            documento.Close(Wd.wdDoNotSaveChanges)
        except Exception:  # noqa: BLE001
            pass


def _guardar_como(documento: Any, salida: Path, formato: int) -> None:
    """Guarda usando SaveAs2, con reserva a SaveAs en versiones antiguas."""
    try:
        documento.SaveAs2(FileName=str(salida), FileFormat=formato)
    except AttributeError:
        documento.SaveAs(FileName=str(salida), FileFormat=formato)


# ---------------------------------------------------------------------------
# 2. Extracción de comentarios
# ---------------------------------------------------------------------------


def extraer_comentarios(sesion: SesionWord, ruta: Path | str) -> list[Comentario]:
    """Extrae los comentarios de un documento con su página real.

    Args:
        sesion: Sesión de Word activa.
        ruta: Documento del que se extraen los comentarios.

    Returns:
        Lista de comentarios en el orden en que aparecen en el documento.

    Raises:
        ErrorDocumento: Si el documento no se puede abrir.
    """
    documento_ruta = _ruta_absoluta(ruta)
    comentarios: list[Comentario] = []

    with sesion.abrir(documento_ruta) as doc:
        _repaginar(doc)
        total = int(doc.Comments.Count)
        registro.info("%s: %d comentarios localizados.", documento_ruta.name, total)
        for indice in range(1, total + 1):
            try:
                comentarios.append(
                    _leer_comentario(doc.Comments.Item(indice), indice, documento_ruta.name)
                )
            except Exception as exc:  # noqa: BLE001
                registro.warning(
                    "Comentario %d de %s ilegible: %s", indice, documento_ruta.name, exc
                )
    return comentarios


def _leer_comentario(comentario: Any, indice: int, nombre_documento: str) -> Comentario:
    """Traduce un objeto Comment de Word a una estructura de datos."""
    ambito = comentario.Scope   # Range del texto comentado, en el cuerpo
    cuerpo = comentario.Range   # Range del texto del propio comentario
    return Comentario(
        documento=nombre_documento,
        indice=indice,
        autor=str(_atributo_seguro(comentario, "Author")),
        iniciales=str(_atributo_seguro(comentario, "Initial")),
        fecha=_texto_fecha(_atributo_seguro(comentario, "Date", None)),
        pagina=_pagina_de(ambito),
        apartado=_apartado_de(ambito),
        texto_comentario=_texto_limpio(cuerpo.Text),
        texto_comentado=_texto_limpio(ambito.Text, maximo=500),
        resuelto="Sí" if _atributo_seguro(comentario, "Done", False) else "No",
    )


def _repaginar(documento: Any) -> None:
    """Fuerza el recálculo de la paginación antes de leer números de página."""
    try:
        documento.Repaginate()
    except Exception:  # noqa: BLE001
        registro.debug("No ha sido posible repaginar el documento.")


# ---------------------------------------------------------------------------
# 3. Extracción de cambios registrados
# ---------------------------------------------------------------------------


def extraer_revisiones(
    sesion: SesionWord,
    ruta: Path | str,
    maximo_caracteres: int = 300,
) -> list[Revision]:
    """Extrae los cambios registrados (control de cambios) de un documento.

    Sirve tanto para un documento con control de cambios del proyectista como
    para el .docx generado por :func:`comparar_documentos`: en ese caso cada
    revisión es una diferencia detectada entre las dos versiones.

    Args:
        sesion: Sesión de Word activa.
        ruta: Documento del que se extraen las revisiones.
        maximo_caracteres: Longitud máxima del texto afectado que se registra.

    Returns:
        Lista de revisiones en orden de aparición.

    Raises:
        ErrorDocumento: Si el documento no se puede abrir.
    """
    documento_ruta = _ruta_absoluta(ruta)
    revisiones: list[Revision] = []

    with sesion.abrir(documento_ruta) as doc:
        _repaginar(doc)
        total = int(doc.Revisions.Count)
        registro.info("%s: %d cambios registrados.", documento_ruta.name, total)
        for indice in range(1, total + 1):
            try:
                revisiones.append(
                    _leer_revision(
                        doc.Revisions.Item(indice), indice,
                        documento_ruta.name, maximo_caracteres,
                    )
                )
            except Exception as exc:  # noqa: BLE001
                registro.warning(
                    "Revisión %d de %s ilegible: %s", indice, documento_ruta.name, exc
                )
    return revisiones


def _leer_revision(
    revision: Any, indice: int, nombre_documento: str, maximo: int
) -> Revision:
    """Traduce un objeto Revision de Word a una estructura de datos."""
    tipo = int(_atributo_seguro(revision, "Type", 0))
    rango = revision.Range
    return Revision(
        documento=nombre_documento,
        indice=indice,
        autor=str(_atributo_seguro(revision, "Author")),
        fecha=_texto_fecha(_atributo_seguro(revision, "Date", None)),
        tipo=tipo,
        descripcion_tipo=TIPOS_DE_REVISION.get(tipo, f"Tipo {tipo}"),
        pagina=_pagina_de(rango),
        apartado=_apartado_de(rango),
        estilo=_estilo_de(rango),
        texto=_texto_limpio(rango.Text, maximo=maximo),
    )


# ---------------------------------------------------------------------------
# 4. Actualización de campos, índices y referencias cruzadas
# ---------------------------------------------------------------------------


def actualizar_campos(
    sesion: SesionWord,
    ruta: Path | str,
    ruta_salida: Path | str,
) -> Path:
    """Actualiza campos, índices y referencias cruzadas, y guarda una copia.

    Recorre todos los "story ranges" del documento (cuerpo, encabezados, pies,
    notas, cuadros de texto), no sólo el cuerpo: las referencias cruzadas y los
    campos de los encabezados quedarían si no sin actualizar.

    El original no se toca: se trabaja sobre una copia y se guarda en
    `ruta_salida`.

    Args:
        sesion: Sesión de Word activa.
        ruta: Documento de origen.
        ruta_salida: Ruta del documento actualizado que se va a generar.

    Returns:
        La ruta del documento generado.

    Raises:
        ErrorDocumento: Si el documento no se puede abrir o guardar.
    """
    origen = _ruta_absoluta(ruta)
    salida = _ruta_absoluta(ruta_salida)
    salida.parent.mkdir(parents=True, exist_ok=True)

    with _carpeta_temporal() as temporal:
        copia = _copia_de_trabajo(origen, temporal)
        with sesion.abrir(copia, solo_lectura=False) as doc:
            # Si el control de cambios estuviera activo, la actualización de
            # campos generaría revisiones espurias en el documento de salida.
            _desactivar_control_de_cambios(doc)
            _actualizar_todos_los_campos(doc)
            _actualizar_indices(doc)
            _guardar_como(doc, salida, Wd.wdFormatDocumentDefault)

    registro.info("Documento actualizado: %s", salida.name)
    return salida


def _desactivar_control_de_cambios(documento: Any) -> None:
    """Desactiva el registro de cambios en la copia de trabajo."""
    try:
        documento.TrackRevisions = False
    except Exception:  # noqa: BLE001
        pass


def _actualizar_todos_los_campos(documento: Any) -> None:
    """Actualiza los campos de todas las historias del documento."""
    for historia in documento.StoryRanges:
        rango = historia
        while rango is not None:
            try:
                rango.Fields.Update()
            except Exception as exc:  # noqa: BLE001
                registro.debug("Campos no actualizados en una historia: %s", exc)
            try:
                rango = rango.NextStoryRange
            except Exception:  # noqa: BLE001
                rango = None


def _actualizar_indices(documento: Any) -> None:
    """Actualiza índices de contenido, de ilustraciones y alfabéticos."""
    colecciones = (
        ("TablesOfContents", "índice de contenido"),
        ("TablesOfFigures", "índice de ilustraciones"),
        ("TablesOfAuthorities", "índice de autoridades"),
        ("Indexes", "índice alfabético"),
    )
    for nombre, descripcion in colecciones:
        coleccion = _atributo_seguro(documento, nombre, None)
        if coleccion is None:
            continue
        try:
            for elemento in coleccion:
                elemento.Update()
        except Exception as exc:  # noqa: BLE001
            registro.debug("No se ha podido actualizar el %s: %s", descripcion, exc)


# ---------------------------------------------------------------------------
# 5. Exportación a PDF
# ---------------------------------------------------------------------------


def exportar_a_pdf(
    sesion: SesionWord,
    ruta: Path | str,
    ruta_salida: Path | str,
    con_marcas_de_revision: bool = False,
    marcadores_desde_titulos: bool = True,
) -> Path:
    """Exporta un documento a PDF con la paginación real de Word.

    Args:
        sesion: Sesión de Word activa.
        ruta: Documento de origen.
        ruta_salida: Ruta del PDF que se va a generar.
        con_marcas_de_revision: Si es True, el PDF muestra el control de
            cambios y los comentarios; si es False, el documento limpio.
        marcadores_desde_titulos: Generar marcadores de PDF a partir de los
            estilos de título.

    Returns:
        La ruta del PDF generado.

    Raises:
        ErrorDocumento: Si el documento no se puede abrir o exportar.
    """
    origen = _ruta_absoluta(ruta)
    salida = _ruta_absoluta(ruta_salida)
    salida.parent.mkdir(parents=True, exist_ok=True)

    with sesion.abrir(origen) as doc:
        try:
            doc.ExportAsFixedFormat(
                OutputFileName=str(salida),
                ExportFormat=Wd.wdExportFormatPDF,
                OpenAfterExport=False,
                OptimizeFor=Wd.wdExportOptimizeForPrint,
                Range=Wd.wdExportAllDocument,
                Item=(
                    Wd.wdExportDocumentWithMarkup
                    if con_marcas_de_revision
                    else Wd.wdExportDocumentContent
                ),
                IncludeDocProps=True,
                KeepIRM=True,
                CreateBookmarks=(
                    Wd.wdExportCreateHeadingBookmarks
                    if marcadores_desde_titulos
                    else Wd.wdExportCreateNoBookmarks
                ),
                DocStructureTags=True,
                BitmapMissingFonts=True,
                UseISO19005_1=False,
            )
        except Exception as exc:  # noqa: BLE001
            raise ErrorDocumento(origen, f"no se ha podido exportar a PDF ({exc})") from exc

    registro.info("PDF generado: %s", salida.name)
    return salida


# ---------------------------------------------------------------------------
# Exportación de tablas a Excel o CSV
# ---------------------------------------------------------------------------


def exportar_tabla(
    filas: Sequence[Any],
    ruta_salida: Path | str,
    titulo_hoja: str = "Datos",
) -> Path:
    """Exporta una lista de dataclasses a Excel (.xlsx) o CSV.

    Si la extensión de `ruta_salida` es .xlsx y openpyxl está disponible, se
    genera una hoja con encabezados, filtro automático y anchos ajustados. En
    caso contrario se genera un CSV con separador ';' y codificación UTF-8 con
    BOM, que Excel en español abre correctamente con doble clic.

    Args:
        filas: Secuencia de dataclasses homogéneas (Comentario o Revision).
        ruta_salida: Ruta del fichero de salida.
        titulo_hoja: Nombre de la hoja de cálculo.

    Returns:
        La ruta del fichero realmente generado.
    """
    salida = _ruta_absoluta(ruta_salida)
    salida.parent.mkdir(parents=True, exist_ok=True)

    if not filas:
        registro.info("Sin datos que exportar a %s.", salida.name)
        return salida

    encabezados = [campo.name for campo in campos_dataclass(filas[0])]
    datos = [[asdict(fila)[clave] for clave in encabezados] for fila in filas]

    if salida.suffix.lower() == ".xlsx":
        try:
            return _exportar_xlsx(encabezados, datos, salida, titulo_hoja)
        except ImportError:
            registro.warning(
                "openpyxl no está instalado; se exporta a CSV. "
                "Instálelo con: pip install openpyxl"
            )
            salida = salida.with_suffix(".csv")
    return _exportar_csv(encabezados, datos, salida)


def _exportar_xlsx(
    encabezados: list[str], datos: list[list[Any]], salida: Path, titulo_hoja: str
) -> Path:
    """Genera un libro de Excel con formato básico de tabla."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo_hoja[:31]
    hoja.append([cabecera.replace("_", " ").capitalize() for cabecera in encabezados])

    relleno = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")

    for fila in datos:
        hoja.append(fila)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = (
        f"A1:{get_column_letter(len(encabezados))}{len(datos) + 1}"
    )
    _ajustar_anchos(hoja, encabezados, datos, get_column_letter)

    libro.save(str(salida))
    registro.info("Fichero Excel generado: %s (%d filas).", salida.name, len(datos))
    return salida


def _ajustar_anchos(
    hoja: Any, encabezados: list[str], datos: list[list[Any]], letra_columna: Any
) -> None:
    """Ajusta el ancho de cada columna al contenido, con un tope razonable."""
    for numero, encabezado in enumerate(encabezados, start=1):
        largo_maximo = max(
            [len(str(encabezado))]
            + [len(str(fila[numero - 1])) for fila in datos[:200]]
        )
        hoja.column_dimensions[letra_columna(numero)].width = min(max(largo_maximo + 2, 10), 70)


def _exportar_csv(encabezados: list[str], datos: list[list[Any]], salida: Path) -> Path:
    """Genera un CSV compatible con Excel en configuración regional española."""
    with salida.open("w", encoding="utf-8-sig", newline="") as fichero:
        escritor = csv.writer(fichero, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        escritor.writerow(encabezados)
        escritor.writerows(datos)
    registro.info("Fichero CSV generado: %s (%d filas).", salida.name, len(datos))
    return salida


# ---------------------------------------------------------------------------
# 6. Procesos en lote
# ---------------------------------------------------------------------------


def emparejar_versiones(
    carpeta: Path | str,
    solo_ultimo_par: bool = True,
) -> list[tuple[Path, Path]]:
    """Empareja versiones sucesivas de documentos por su sufijo _vNN.

    Agrupa los ficheros por nombre base (todo lo anterior al sufijo de versión)
    y los ordena por número de versión. Ejemplo::

        PROY_Señalizacion_v03.docx
        PROY_Señalizacion_v04.docx   ->  (v03, v04)

    Args:
        carpeta: Carpeta donde buscar los documentos.
        solo_ultimo_par: Si es True devuelve sólo el par más reciente de cada
            documento; si es False, todos los pares consecutivos.

    Returns:
        Lista de tuplas (versión anterior, versión posterior).
    """
    raiz = _ruta_absoluta(carpeta)
    grupos: dict[str, list[tuple[int, Path]]] = {}

    for fichero in sorted(raiz.iterdir()):
        if not fichero.is_file() or not _es_documento_word(fichero):
            continue
        coincidencia = PATRON_VERSION.match(fichero.stem)
        if not coincidencia:
            registro.debug("Sin sufijo de versión, se omite: %s", fichero.name)
            continue
        base = coincidencia.group("base")
        grupos.setdefault(base, []).append((int(coincidencia.group("num")), fichero))

    pares: list[tuple[Path, Path]] = []
    for base, versiones in sorted(grupos.items()):
        if len(versiones) < 2:
            registro.warning("Sólo hay una versión de '%s'; no se puede comparar.", base)
            continue
        versiones.sort(key=lambda elemento: elemento[0])
        consecutivos = [
            (versiones[i][1], versiones[i + 1][1]) for i in range(len(versiones) - 1)
        ]
        pares.extend(consecutivos[-1:] if solo_ultimo_par else consecutivos)

    registro.info("Pares de versiones detectados: %d", len(pares))
    return pares


def comparar_lote(
    carpeta_entrada: Path | str,
    carpeta_salida: Path | str,
    solo_ultimo_par: bool = True,
    exportar_hallazgos: bool = True,
    autor_revision: str = "Comparación automática",
) -> ResultadoLote:
    """Compara todos los pares de versiones de una carpeta.

    Por cada par genera un documento "COMPARATIVA_<base>_vNN_vs_vMM.docx" y,
    opcionalmente, un Excel con el listado de diferencias detectadas, que es la
    base del registro de hallazgos.

    Un documento protegido, bloqueado o dañado se registra como incidencia y no
    interrumpe el procesamiento de los restantes.

    Args:
        carpeta_entrada: Carpeta con las versiones a comparar.
        carpeta_salida: Carpeta donde se depositan los resultados.
        solo_ultimo_par: Comparar sólo las dos versiones más recientes.
        exportar_hallazgos: Generar además el Excel de diferencias.
        autor_revision: Nombre que figurará como autor de las marcas.

    Returns:
        Resumen de la ejecución.
    """
    salida = _ruta_absoluta(carpeta_salida)
    salida.mkdir(parents=True, exist_ok=True)
    pares = emparejar_versiones(carpeta_entrada, solo_ultimo_par)
    resultado = ResultadoLote(correctos=[], fallidos=[], salidas=[])

    with SesionWord() as sesion:
        for anterior, nueva in pares:
            etiqueta = f"{anterior.name} -> {nueva.name}"
            try:
                ficheros = _comparar_par(
                    sesion, anterior, nueva, salida, exportar_hallazgos, autor_revision
                )
                resultado.correctos.append(etiqueta)
                resultado.salidas.extend(ficheros)
            except ErrorDocumento as exc:
                registro.error("INCIDENCIA en %s: %s", etiqueta, exc.motivo)
                resultado.fallidos.append((etiqueta, exc.motivo))
            except Exception as exc:  # noqa: BLE001
                registro.exception("Error inesperado en %s", etiqueta)
                resultado.fallidos.append((etiqueta, str(exc)))

    registro.info("Comparación en lote finalizada.\n%s", resultado.resumen())
    return resultado


def _comparar_par(
    sesion: SesionWord,
    anterior: Path,
    nueva: Path,
    carpeta_salida: Path,
    exportar_hallazgos: bool,
    autor_revision: str,
) -> list[Path]:
    """Compara un par de versiones y genera los ficheros de salida."""
    nombre = _nombre_comparativa(anterior, nueva)
    documento_comparado = comparar_documentos(
        sesion, anterior, nueva, carpeta_salida / f"{nombre}.docx",
        autor_revision=autor_revision,
    )
    generados = [documento_comparado]

    if exportar_hallazgos:
        diferencias = extraer_revisiones(sesion, documento_comparado)
        generados.append(
            exportar_tabla(
                diferencias,
                carpeta_salida / f"{nombre}_DIFERENCIAS.xlsx",
                titulo_hoja="Diferencias",
            )
        )
    return generados


def _nombre_comparativa(anterior: Path, nueva: Path) -> str:
    """Construye el nombre base del fichero comparativo."""
    coincidencia_a = PATRON_VERSION.match(anterior.stem)
    coincidencia_b = PATRON_VERSION.match(nueva.stem)
    if coincidencia_a and coincidencia_b:
        base = coincidencia_a.group("base")
        return (
            f"COMPARATIVA_{base}_v{coincidencia_a.group('num')}"
            f"_vs_v{coincidencia_b.group('num')}"
        )
    return f"COMPARATIVA_{anterior.stem}_vs_{nueva.stem}"


def extraer_comentarios_lote(
    carpeta_entrada: Path | str,
    carpeta_salida: Path | str,
    fichero_unificado: bool = True,
    incluir_revisiones: bool = False,
) -> ResultadoLote:
    """Extrae los comentarios de todos los documentos de una carpeta.

    Args:
        carpeta_entrada: Carpeta con los documentos a procesar.
        carpeta_salida: Carpeta donde se depositan los ficheros generados.
        fichero_unificado: Si es True genera un único Excel con todos los
            documentos; si es False, uno por documento.
        incluir_revisiones: Generar además la tabla de cambios registrados.

    Returns:
        Resumen de la ejecución.
    """
    entrada = _ruta_absoluta(carpeta_entrada)
    salida = _ruta_absoluta(carpeta_salida)
    salida.mkdir(parents=True, exist_ok=True)

    documentos = [f for f in sorted(entrada.iterdir())
                  if f.is_file() and _es_documento_word(f)]
    registro.info("Documentos a procesar: %d", len(documentos))

    resultado = ResultadoLote(correctos=[], fallidos=[], salidas=[])
    todos_comentarios: list[Comentario] = []
    todas_revisiones: list[Revision] = []

    with SesionWord() as sesion:
        for documento in documentos:
            try:
                comentarios = extraer_comentarios(sesion, documento)
                revisiones = (
                    extraer_revisiones(sesion, documento) if incluir_revisiones else []
                )
                todos_comentarios.extend(comentarios)
                todas_revisiones.extend(revisiones)
                if not fichero_unificado:
                    resultado.salidas.extend(
                        _exportar_por_documento(documento, salida, comentarios, revisiones)
                    )
                resultado.correctos.append(documento.name)
            except ErrorDocumento as exc:
                registro.error("INCIDENCIA en %s: %s", documento.name, exc.motivo)
                resultado.fallidos.append((documento.name, exc.motivo))
            except Exception as exc:  # noqa: BLE001
                registro.exception("Error inesperado en %s", documento.name)
                resultado.fallidos.append((documento.name, str(exc)))

    if fichero_unificado:
        resultado.salidas.extend(
            _exportar_unificado(salida, todos_comentarios, todas_revisiones)
        )

    registro.info("Extracción en lote finalizada.\n%s", resultado.resumen())
    return resultado


def _exportar_por_documento(
    documento: Path,
    carpeta_salida: Path,
    comentarios: list[Comentario],
    revisiones: list[Revision],
) -> list[Path]:
    """Genera un fichero de salida por cada documento procesado."""
    generados = [
        exportar_tabla(
            comentarios,
            carpeta_salida / f"COMENTARIOS_{documento.stem}.xlsx",
            titulo_hoja="Comentarios",
        )
    ]
    if revisiones:
        generados.append(
            exportar_tabla(
                revisiones,
                carpeta_salida / f"CAMBIOS_{documento.stem}.xlsx",
                titulo_hoja="Cambios registrados",
            )
        )
    return generados


def _exportar_unificado(
    carpeta_salida: Path,
    comentarios: list[Comentario],
    revisiones: list[Revision],
) -> list[Path]:
    """Genera un único fichero con los datos de todos los documentos."""
    marca = datetime.now().strftime("%Y%m%d_%H%M")
    generados = [
        exportar_tabla(
            comentarios,
            carpeta_salida / f"COMENTARIOS_CONSOLIDADO_{marca}.xlsx",
            titulo_hoja="Comentarios",
        )
    ]
    if revisiones:
        generados.append(
            exportar_tabla(
                revisiones,
                carpeta_salida / f"CAMBIOS_CONSOLIDADO_{marca}.xlsx",
                titulo_hoja="Cambios registrados",
            )
        )
    return generados


# ---------------------------------------------------------------------------
# Diagnóstico
# ---------------------------------------------------------------------------


def diagnostico() -> bool:
    """Comprueba que el entorno puede pilotar Word por COM.

    Arranca Word, muestra la versión instalada y lo cierra. Conviene ejecutarlo
    la primera vez, antes de procesar documentos reales.

    Returns:
        True si Word ha respondido correctamente.
    """
    configurar_registro()
    registro.info("Iniciando comprobación del entorno…")
    try:
        with SesionWord() as sesion:
            app = sesion.app
            registro.info("Versión de Word: %s", _atributo_seguro(app, "Version"))
            registro.info("Compilación: %s", _atributo_seguro(app, "Build"))
            registro.info("Ruta del ejecutable: %s", _atributo_seguro(app, "Path"))
            registro.info("Documentos abiertos en esta instancia: %s",
                          app.Documents.Count)
        registro.info("Comprobación superada: el entorno está listo.")
        return True
    except Exception as exc:  # noqa: BLE001
        registro.error("Comprobación fallida: %s", exc)
        registro.error(
            "Verifique que Microsoft Word está instalado y que pywin32 lo está "
            "también (pip install pywin32)."
        )
        return False


if __name__ == "__main__":
    diagnostico()
