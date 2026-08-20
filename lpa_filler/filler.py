"""Preenche o template .xlsm de LPA a partir dos dados do projeto.

Regenera as regiões de dados (Portada, Control de versiones, Doc Evaluados e LPA)
clonando o estilo de linhas-modelo do template, replicando as mesclas e as
fórmulas, e preservando macros/dropdowns/imagens via ``extensions.preserve``.
"""
from __future__ import annotations

import warnings
from pathlib import Path
from typing import Any

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
from openpyxl.utils import column_index_from_string as col_idx

from openpyxl.styles import Font

from . import extensions, model, novidade, styles

# ----- Mapa de colunas por aba (índice 1-based) -----
CV = {"rev": col_idx("A"), "fecha": col_idx("B"), "descripcion": col_idx("C")}
DE = {c: col_idx(c) for c in "ABCDEFGHIJ"}
LPA = {c: col_idx(c) for c in "ABCDEFGHIJK"}

# Azul da Exceltic para conteúdo novo desde o LPA anterior (pedido do avaliador,
# não do PE/05 — é convenção da casa, não normativa). ARGB com alpha opaco.
COR_NOVO = "FF0070C0"

# Linhas a partir das quais começam os dados (a seguir aos cabeçalhos).
CV_FIRST, DE_FIRST, LPA_FIRST = 4, 3, 2

# Fórmulas do template. O intervalo do VLOOKUP cobre até à última linha
# realmente gerada em Doc Evaluados (o template trazia $G$148 fixo, que
# quebrava silenciosamente com mais de ~146 linhas de documentos).
VLOOKUP_REF = "=VLOOKUP(C{r},'Doc Evaluados'!$A$1:$G${last},2,FALSE)"
DE_ESTADO_FORMULA = '=IF(COUNTIFS(LPA!J:J,"Abierto",LPA!C:C,{ref})>0,"Abierto","Cerrado")'
DE_COMENT_FORMULA = '=IF({estado}="Abierto","Ver pestaña LPA con los hallazgos identificados","Versión formal sin hallazgos")'
LPA_J_FIRST = "=IF(K{r}<>0,K{r},#REF!)"
LPA_J_NEXT = "=IF(K{r}<>0,K{r},J{p})"


def fill(
    template: str | Path,
    data: dict[str, Any],
    output: str | Path,
    veredicto_text: str | None = None,
    veredicto_cell: str | None = None,
    skip_lpa: bool = False,
    anterior: dict[str, Any] | None = None,
) -> Path:
    """Gera o .xlsm. ``veredicto_text`` fica registado nas propriedades do
    ficheiro (Ficheiro > Informações no Excel) e, se ``veredicto_cell`` for
    indicado ("Aba!Célula", ex. "Portada!B30"), também nessa célula — a célula
    não tem default porque a posição livre depende do template de cada obra.
    Se ``skip_lpa`` for True, deixa a aba LPA vazia para editar manualmente.

    ``anterior`` (o LPA já emitido, na mesma forma de ``data`` — ver
    ``extract.extract``) faz o que for novo desde ele sair a azul (#0070C0):
    puntos inteiros, diálogo acrescentado a puntos que já existiam, documentos/
    envíos novos, a revisão nova. Sem ``anterior`` (primeira vez), nada se
    destaca — não há com que comparar."""
    template, output = Path(template), Path(output)
    wb = openpyxl.load_workbook(template, keep_vba=True)
    nov = novidade.calcular(data, anterior)

    # Capturar estilos das linhas-modelo ANTES de limpar as regiões.
    cv_style = styles.snapshot_row_styles(wb["Control de versiones"], CV_FIRST, 3)
    de_first = styles.snapshot_row_styles(wb["Doc Evaluados"], DE_FIRST, 10)
    de_sub = styles.snapshot_row_styles(wb["Doc Evaluados"], DE_FIRST + 1, 10)
    lpa_first = styles.snapshot_row_styles(wb["LPA"], LPA_FIRST, 11)
    lpa_resp = styles.snapshot_row_styles(wb["LPA"], LPA_FIRST + 1, 11)

    _fill_portada(wb["Portada"], data.get("portada", {}))
    _fill_versiones(wb["Control de versiones"], data.get("versiones", []), cv_style,
                     nov["versiones_novas"])
    de_last = _fill_documentos(wb["Doc Evaluados"], data.get("documentos", []), de_first, de_sub,
                                nov["documentos_novos"], nov["envios_novos"])
    if skip_lpa:
        # Limpar a aba LPA deixando espaço vazio para editar manualmente
        styles.clear_region(wb["LPA"], LPA_FIRST, max(wb["LPA"].max_row, LPA_FIRST), 11)
        lpa_last = LPA_FIRST - 1
    else:
        lpa_last = _fill_lpa(wb["LPA"], data.get("puntos", []), lpa_first, lpa_resp, de_last,
                              nov["puntos_novos"], nov["dialogos_novos"])

    if veredicto_text:
        wb.properties.description = veredicto_text
        if veredicto_cell and "!" in veredicto_cell:
            sheet_name, cell = veredicto_cell.split("!", 1)
            wb[sheet_name][cell.strip()] = veredicto_text

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    # Repor o que o openpyxl descarta e expandir as validações às linhas geradas.
    row_overrides = {"Doc Evaluados": de_last}
    if not skip_lpa:
        row_overrides["LPA"] = lpa_last
    extensions.preserve(output, template, row_overrides=row_overrides)
    return output


# --------------------------------------------------------------------------- #
def _anchor(ws, row, col):
    """Se (row, col) estiver dentro de uma mescla, devolve a célula-âncora
    (canto superior esquerdo), que é a única onde se pode escrever."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return rng.min_row, rng.min_col
    return row, col


def _set(ws, row, col, value):
    if value is None:
        return
    r, c = _anchor(ws, row, col)
    ws.cell(row=r, column=c).value = value


def _marcar_novo(ws, row, col) -> None:
    """Pinta o texto desta célula a azul (COR_NOVO) — conteúdo novo desde o LPA
    anterior. Só troca a cor da fonte; tamanho, negrito e tipo de letra do
    template mantêm-se, porque se clona o Font existente e só se muda ``color``."""
    r, c = _anchor(ws, row, col)
    cel = ws.cell(row=r, column=c)
    f = cel.font
    cel.font = Font(
        name=f.name, size=f.size, bold=f.bold, italic=f.italic,
        underline=f.underline, strike=f.strike, vertAlign=f.vertAlign,
        color=COR_NOVO,
    )


def _fill_portada(ws, p: dict[str, Any]) -> None:
    _set(ws, 12, col_idx("B"), p.get("titulo"))
    _set(ws, 14, col_idx("B"), p.get("subtitulo"))
    _set(ws, 16, col_idx("B"), p.get("referencia"))
    _set(ws, 20, col_idx("B"), p.get("normativa"))
    # Até 3 evaluadores nas colunas B, C, E (linha 26 nome, 27 papel).
    cols = [col_idx("B"), col_idx("C"), col_idx("E")]
    for i, ev in enumerate(p.get("evaluadores", [])[:3]):
        _set(ws, 26, cols[i], ev.get("nombre"))
        _set(ws, 27, cols[i], ev.get("rol"))


def _fill_versiones(ws, versiones: list[dict], style: dict, versiones_novas: set) -> None:
    styles.clear_region(ws, CV_FIRST, max(ws.max_row, CV_FIRST), 3)
    for i, v in enumerate(versiones):
        r = CV_FIRST + i
        styles.apply_row_styles(ws, r, style)
        _set(ws, r, CV["rev"], v.get("rev", i + 1))
        _set(ws, r, CV["fecha"], v.get("fecha"))
        _set(ws, r, CV["descripcion"], v.get("descripcion"))
        if v.get("rev") in versiones_novas:
            for col in CV.values():
                _marcar_novo(ws, r, col)


def _fill_documentos(ws, documentos: list[dict], first_style: dict, sub_style: dict,
                      documentos_novos: set, envios_novos: dict[str, set]) -> int:
    styles.clear_region(ws, DE_FIRST, max(ws.max_row, DE_FIRST), 10)
    r = DE_FIRST
    spans: list[tuple] = []  # (categoria, start, end)

    for doc in documentos:
        cat = doc.get("categoria")
        envios = doc.get("envios") or [{}]
        nome = doc.get("nombre")
        doc_novo = nome in documentos_novos
        idx_novos = set() if doc_novo else envios_novos.get(nome, set())
        start = r
        for k, env in enumerate(envios):
            row = r + k
            styles.apply_row_styles(ws, row, first_style if k == 0 else sub_style)
            styles.copy_row_height(ws, DE_FIRST + (0 if k == 0 else 1), row)
            _set(ws, row, DE["C"], env.get("version"))
            _set(ws, row, DE["D"], env.get("fecha"))
            _set(ws, row, DE["E"], env.get("autor"))
            _set(ws, row, DE["F"], env.get("envio"))
            _set(ws, row, DE["G"], env.get("fecha_envio"))
            if not cat:
                _set(ws, row, DE["B"], env.get("referencia"))
            if k in idx_novos:
                # Só este envío é novo: o documento já existia, não se pinta
                # a coluna B/H/I/J, que descrevem o documento inteiro.
                cols = ("C", "D", "E", "F", "G") if cat else ("B", "C", "D", "E", "F", "G")
                for c in cols:
                    _marcar_novo(ws, row, DE[c])
        end = r + len(envios) - 1

        if cat:
            # Com categoria: B = nome do documento (mesclado); A = categoria (passo 2).
            _set(ws, start, DE["B"], doc.get("nombre"))
            styles.merge(ws, DE["B"], start, end)
            name_cell = f"B{start}"
        else:
            _set(ws, start, DE["A"], doc.get("nombre"))
            styles.merge(ws, DE["A"], start, end)
            name_cell = f"A{start}"

        # Firmado / Estado / Comentarios (mesclados no documento).
        _set(ws, start, DE["H"], doc.get("firmado", "NA"))
        estado = doc.get("estado", "auto")
        ws.cell(row=start, column=DE["I"]).value = (
            DE_ESTADO_FORMULA.format(ref=name_cell) if estado in (None, "auto") else estado
        )
        ws.cell(row=start, column=DE["J"]).value = DE_COMENT_FORMULA.format(estado=f"I{start}")
        for c in ("H", "I", "J"):
            styles.merge(ws, DE[c], start, end)

        if doc_novo:
            # Documento inteiro novo: todas as linhas, todas as colunas.
            for row in range(start, end + 1):
                for c in DE.values():
                    _marcar_novo(ws, row, c)

        spans.append((cat, start, end))
        r = end + 1

    # Passo 2: mesclar a coluna A por corridas consecutivas da mesma categoria.
    i = 0
    while i < len(spans):
        cat, s, _ = spans[i]
        if cat:
            j = i
            while j + 1 < len(spans) and spans[j + 1][0] == cat:
                j += 1
            _set(ws, s, DE["A"], cat)
            styles.merge(ws, DE["A"], s, spans[j][2])
            i = j + 1
        else:
            i += 1
    return r - 1


def _fill_lpa(ws, puntos: list[dict], first_style: dict, resp_style: dict, de_last: int,
              puntos_novos: set, dialogos_novos: dict[str, set]) -> int:
    styles.clear_region(ws, LPA_FIRST, max(ws.max_row, LPA_FIRST), 11)
    r = LPA_FIRST
    for pt in puntos:
        punto_novo = pt.get("id") in puntos_novos
        idx_novos = set() if punto_novo else dialogos_novos.get(pt.get("id"), set())
        dialogo = pt.get("dialogo") or [{}]
        start = r
        for k, d in enumerate(dialogo):
            row = r + k
            styles.apply_row_styles(ws, row, first_style if k == 0 else resp_style)
            styles.copy_row_height(ws, LPA_FIRST + (0 if k == 0 else 1), row)
            _set(ws, row, LPA["G"], d.get("tipo"))
            _set(ws, row, LPA["I"], d.get("texto"))
            ws.cell(row=row, column=LPA["J"]).value = (
                LPA_J_FIRST.format(r=row) if k == 0 else LPA_J_NEXT.format(r=row, p=row - 1)
            )
            if punto_novo or k in idx_novos:
                # Só tipo+texto: a coluna J é fórmula (arrasta o estado para baixo),
                # não conteúdo escrito por alguém — não faz sentido pintá-la.
                _marcar_novo(ws, row, LPA["G"])
                _marcar_novo(ws, row, LPA["I"])
        end = r + len(dialogo) - 1

        _set(ws, start, LPA["A"], pt.get("n"))
        _set(ws, start, LPA["B"], pt.get("eval"))
        _set(ws, start, LPA["C"], pt.get("documento"))
        ref = pt.get("ref_documento", "auto")
        ws.cell(row=start, column=LPA["D"]).value = (
            VLOOKUP_REF.format(r=start, last=max(de_last, DE_FIRST)) if ref in (None, "auto") else ref
        )
        _set(ws, start, LPA["E"], pt.get("punto"))
        _set(ws, start, LPA["F"], pt.get("valoracion"))
        _set(ws, start, LPA["H"], pt.get("version"))
        _set(ws, start, LPA["K"], pt.get("estado"))
        for c in ("A", "B", "C", "D", "E", "F", "K"):
            styles.merge(ws, LPA[c], start, end)
            if punto_novo:
                _marcar_novo(ws, start, LPA[c])
        r = end + 1
    return r - 1
