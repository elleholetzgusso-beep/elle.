"""Exporta a base de hallazgos como um .xlsx organizado — guia de "por onde começar".

Ao arrancar uma obra nova, o avaliador quer ver o que já apareceu em obras
anteriores para saber onde focar. Este módulo lê a base do ``harvest`` e escreve
um Excel com duas vistas:

  1. "Por onde começar": contagens por TIPO de documento e por TEMA (RAMS),
     ordenadas pela gravidade (Críticos primeiro) — mostra onde costumam surgir
     hallazgos e quão graves são.
  2. "Hallazgos": todos os achados, ordenados e filtráveis, com as colunas úteis,
     cabeçalho fixo e cor por valoración.

Determinístico. O tipo de documento sai do nome por regras (como o ``tema`` sai
do texto); nada é inventado.
"""
from __future__ import annotations

import csv
import re
import unicodedata
from pathlib import Path
from typing import Any

from . import tema as _tema

# Tipo de documento a partir do nome (ordem = prioridade). Generaliza entre obras.
TIPOS: list[tuple[str, str]] = [
    ("Safety Case / Caso de Seguridad", r"safety case|caso de seguridad|dossier de seguridad"),
    # F3 (DS-REP = Definición del Sistema + Registro de Peligros) vem ANTES do REP:
    # é um documento combinado cujo nome contém 'REP', mas a identidade é o F3.
    ("Definición del Sistema (F3)", r"definicion de(l)? sistema|\bf3\b|ds-rep"),
    ("REP / Registro de Peligros", r"\brep\b|registro.{0,3} de peligros|hazard log|registro especifico de peligros"),
    ("Estudio Previo de Seguridad", r"estudio previo.{0,4} de seguridad|estudio previo seguridad"),
    ("Informe de Validación", r"informe de validacion|\bival\b"),
    ("Informe de No Regresión", r"no regresion|\binr\b"),
    ("Informe de Seguridad", r"informe de seguridad"),
    ("Plan de Pruebas / PeS", r"plan de pruebas|puesta en servicio|\bpes\b|\bpps\b|registro.{0,3}pruebas|acta.{0,3}pruebas"),
    ("Plan de Seguridad", r"plan de seguridad|plan de gestion de la seguridad"),
    ("SRAC", r"\bsrac\b|condiciones de aplicacion"),
    ("Pliego (PPTP/PCAP)", r"\bpptp\b|\bpcap\b|pliego"),
    ("Software / Aplicación", r"application preparation|product version|aplicacion especifica|\bsoftware\b|ebilock"),
    ("Firmas / Documentación", r"firmas|listado de documentacion|documentacion evaluada"),
    ("General", r"^general$"),
]

VALORACIONES = ["Crítico", "Importante", "Informativo", "Formal"]
_VAL_ORDER = {v: i for i, v in enumerate(VALORACIONES)}


def _norm(s: Any) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", str(s or "")) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def tipo_documento(nombre: str) -> str:
    """Classifica o nome de um documento num tipo genérico (entre obras)."""
    n = _norm(nombre)
    if not n:
        return "Otros"
    for rotulo, padrao in TIPOS:
        if re.search(padrao, n):
            return rotulo
    return "Otros"


def _temas_da_linha(row: dict) -> list[str]:
    """Temas RAMS da linha: usa a coluna 'tema' se existir, senão classifica o texto."""
    if (row.get("tema") or "").strip():
        return [t.strip() for t in row["tema"].split(",") if t.strip()]
    texto = " ".join(str(row.get(k) or "") for k in ("hallazgo", "discusion", "documento"))
    return _tema.classify(texto)


def _contar(rows: list[dict], chave) -> list[dict]:
    """Agrega por 'chave(row) -> str|list', contando por valoración + nº de obras."""
    acc: dict[str, dict] = {}
    for r in rows:
        val = (r.get("valoracion") or "").strip()
        obra = (r.get("obra") or "").strip()
        chaves = chave(r)
        if isinstance(chaves, str):
            chaves = [chaves]
        for k in chaves:
            d = acc.setdefault(k, {v: 0 for v in VALORACIONES} | {"Total": 0, "_obras": set()})
            if val in d:
                d[val] += 1
            d["Total"] += 1
            if obra:
                d["_obras"].add(obra)
    saida = []
    for k, d in acc.items():
        saida.append({"grupo": k, **{v: d[v] for v in VALORACIONES},
                      "Total": d["Total"], "Obras": len(d["_obras"])})
    # ordena por Críticos desc, depois Total desc
    saida.sort(key=lambda x: (-x["Crítico"], -x["Total"]))
    return saida


def load_base(csv_path: str | Path) -> list[dict]:
    with Path(csv_path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def exportar(rows: list[dict], out_xlsx: str | Path) -> dict[str, int]:
    """Escreve o .xlsx organizado. Devolve algumas contagens para o resumo na consola."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # Enriquecer cada linha com tipo e temas (uma linha da base = um hallazgo).
    for r in rows:
        r["_tipo"] = tipo_documento(r.get("documento", ""))
        r["_temas"] = _temas_da_linha(r)

    wb = Workbook()

    # Estilos
    accent = "1D4E6E"
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor=accent)
    val_fill = {
        "Crítico": PatternFill("solid", fgColor="F7D9DC"),
        "Importante": PatternFill("solid", fgColor="F7ECD3"),
        "Informativo": PatternFill("solid", fgColor="DEECF6"),
        "Formal": PatternFill("solid", fgColor="E9ECEF"),
    }
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    def _cabecalho(ws, headers):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"

    # ---- Aba 1: Por onde começar ----
    ws1 = wb.active
    ws1.title = "Por onde começar"
    cols = ["Grupo", *VALORACIONES, "Total", "Nº obras"]

    linha = 1
    for titulo, chave in (("POR TIPO DE DOCUMENTO", lambda r: r["_tipo"]),
                          ("POR TEMA (RAMS/CENELEC)", lambda r: r["_temas"] or ["(sem tema)"])):
        c = ws1.cell(row=linha, column=1, value=titulo)
        c.font = Font(bold=True, size=12, color=accent)
        linha += 1
        for i, h in enumerate(cols, 1):
            cell = ws1.cell(row=linha, column=i, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
        linha += 1
        for d in _contar(rows, chave):
            ws1.cell(row=linha, column=1, value=d["grupo"])
            for i, v in enumerate(VALORACIONES, 2):
                cell = ws1.cell(row=linha, column=i, value=d[v])
                if d[v]:
                    cell.fill = val_fill[v]
            ws1.cell(row=linha, column=6, value=d["Total"]).font = Font(bold=True)
            ws1.cell(row=linha, column=7, value=d["Obras"])
            linha += 1
        linha += 2  # espaço entre as duas tabelas
    widths1 = [34, 11, 12, 12, 10, 9, 9]
    for i, w in enumerate(widths1, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ---- Aba 2: Hallazgos (organizados e filtráveis) ----
    ws2 = wb.create_sheet("Hallazgos")
    headers = ["Tipo doc", "Valoración", "Estado", "Tema", "Obra", "Documento",
               "Punto", "Hallazgo", "Discusión", "Fonte (LPA)"]
    _cabecalho(ws2, headers)
    ordenadas = sorted(
        rows,
        key=lambda r: (r["_tipo"], _VAL_ORDER.get((r.get("valoracion") or "").strip(), 9),
                       ", ".join(r["_temas"])),
    )
    for r in ordenadas:
        val = (r.get("valoracion") or "").strip()
        linha_vals = [
            r["_tipo"], val, (r.get("estado") or "").strip(), ", ".join(r["_temas"]),
            (r.get("obra") or "").strip(), (r.get("documento") or "").strip(),
            (r.get("punto") or "").strip(), (r.get("hallazgo") or "").strip(),
            (r.get("discusion") or "").strip(), (r.get("fuente") or "").strip(),
        ]
        ws2.append(linha_vals)
        rr = ws2.max_row
        if val in val_fill:
            ws2.cell(row=rr, column=2).fill = val_fill[val]
        for col in (8, 9):  # Hallazgo, Discusión
            ws2.cell(row=rr, column=col).alignment = wrap
        for col in (1, 2, 3, 4, 5, 6, 7, 10):
            ws2.cell(row=rr, column=col).alignment = top
    widths2 = [26, 12, 10, 22, 18, 30, 20, 60, 40, 22]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws2.max_row}"

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    return {"hallazgos": len(rows),
            "tipos": len({r["_tipo"] for r in rows}),
            "obras": len({(r.get("obra") or "").strip() for r in rows if (r.get("obra") or "").strip()})}
