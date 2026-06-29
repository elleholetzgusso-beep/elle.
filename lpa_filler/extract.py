"""Caminho inverso: lê um .xlsm já preenchido e reconstrói o YAML de dados.

Útil para arrancar de um LPA existente. É "best-effort": o aninhamento de
categorias em Doc Evaluados (ex.: "Planos") é achatado em documentos simples.
"""
from __future__ import annotations

import warnings
from pathlib import Path
from typing import Any

import openpyxl

from .filler import CV_FIRST, DE_FIRST, LPA_FIRST


def _val(ws, r, c):
    import datetime as dt

    v = ws.cell(row=r, column=c).value
    if v == "":
        return None
    if isinstance(v, dt.datetime) and v.time() == dt.time(0, 0):
        return v.date()
    return v


def _block_size(ws, start_row: int, key_col: int, max_row: int) -> int:
    """Quantas linhas pertencem ao bloco que começa em start_row (até ao próximo
    valor não-vazio na coluna-chave)."""
    n = 1
    r = start_row + 1
    while r <= max_row and _val(ws, r, key_col) in (None, ""):
        n += 1
        r += 1
    return n


def extract(xlsm_path: str | Path) -> dict[str, Any]:
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(xlsm_path, keep_vba=True, data_only=False)
    data: dict[str, Any] = {}

    # ---- Portada ----
    p = wb["Portada"]
    portada: dict[str, Any] = {}
    for key, cell in {"titulo": "B12", "subtitulo": "B14", "referencia": "B16", "normativa": "B20"}.items():
        if p[cell].value:
            portada[key] = p[cell].value
    evals = []
    for col in ("B", "C", "E"):
        nombre = p[f"{col}26"].value
        if nombre:
            evals.append({"nombre": nombre, "rol": p[f"{col}27"].value or ""})
    if evals:
        portada["evaluadores"] = evals
    data["portada"] = portada

    # ---- Control de versiones ----
    cv = wb["Control de versiones"]
    versiones = []
    for r in range(CV_FIRST, cv.max_row + 1):
        if _val(cv, r, 1) is None and _val(cv, r, 3) is None:
            continue
        versiones.append(
            {"rev": _val(cv, r, 1), "fecha": _val(cv, r, 2), "descripcion": _val(cv, r, 3)}
        )
    data["versiones"] = versiones

    # ---- Doc Evaluados ----
    de = wb["Doc Evaluados"]
    documentos = []
    r = DE_FIRST
    while r <= de.max_row:
        nombre = _val(de, r, 1) or _val(de, r, 2)
        if nombre is None:
            r += 1
            continue
        size = _block_size(de, r, 1, de.max_row)  # agrupa pela coluna A (nome do documento)
        envios = []
        for k in range(size):
            row = r + k
            envios.append(
                {
                    "referencia": _val(de, row, 2),
                    "version": _val(de, row, 3),
                    "fecha": _val(de, row, 4),
                    "autor": _val(de, row, 5),
                    "envio": _val(de, row, 6),
                    "fecha_envio": _val(de, row, 7),
                }
            )
        estado = _val(de, r, 9)
        if isinstance(estado, str) and estado.startswith("="):
            estado = "auto"
        documentos.append(
            {"nombre": nombre, "firmado": _val(de, r, 8) or "NA", "estado": estado, "envios": envios}
        )
        r += size
    data["documentos"] = documentos

    # ---- LPA ----
    lpa = wb["LPA"]
    puntos = []
    r = LPA_FIRST
    while r <= lpa.max_row:
        if _val(lpa, r, 1) is None:
            r += 1
            continue
        size = _block_size(lpa, r, 1, lpa.max_row)  # blocos pela coluna A (Nº)
        dialogo = []
        for k in range(size):
            row = r + k
            tipo, texto = _val(lpa, row, 7), _val(lpa, row, 9)
            if tipo is not None or texto is not None:
                dialogo.append({"tipo": tipo, "texto": texto})
        ref = _val(lpa, r, 4)
        if isinstance(ref, str) and ref.startswith("="):
            ref = "auto"
        puntos.append(
            {
                "n": _val(lpa, r, 1),
                "eval": _val(lpa, r, 2),
                "documento": _val(lpa, r, 3),
                "ref_documento": ref,
                "punto": _val(lpa, r, 5),
                "valoracion": _val(lpa, r, 6),
                "version": _val(lpa, r, 8),
                "estado": _val(lpa, r, 11),
                "dialogo": dialogo,
            }
        )
        r += size
    data["puntos"] = puntos

    return data
