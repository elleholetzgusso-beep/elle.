"""Testes do exportador da base organizada (guide.py)."""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from lpa_filler import guide  # noqa: E402


def test_tipo_documento_generaliza_entre_obras():
    assert guide.tipo_documento("Anejo 31 - Estudio Previo de Seguridad") == "Estudio Previo de Seguridad"
    assert guide.tipo_documento("Apéndice 1. REP") == "REP / Registro de Peligros"
    assert guide.tipo_documento("Caso de Seguridad") == "Safety Case / Caso de Seguridad"
    assert guide.tipo_documento("Informe de seguridad") == "Informe de Seguridad"
    assert guide.tipo_documento("F3. Definición del Sistema (DS-REP)") == "Definición del Sistema (F3)"
    assert guide.tipo_documento("Plan de Pruebas y PeS Cullera") == "Plan de Pruebas / PeS"
    assert guide.tipo_documento("Documento raríssimo qualquer") == "Otros"


def test_tipo_documento_reconhece_fichas_sag_csm_ra():
    """As fichas F3..F7 trazem quase todas '-REP' no nome sem serem o REP:
    a identidade é o nº da ficha, e é ele que manda (senão tudo cai em REP)."""
    t = guide.tipo_documento
    assert t("SAG1-AP2-F3 DS-REP. DE LA L352 Y ANDEN 3 TP_v0.5") == "Definición del Sistema (F3)"
    assert t("SAG1-AP2-F5 REP-L352 Y ANDEN 3 TP_v1.0_fe") == "REP / Registro de Peligros"
    assert t("SAG1.AP2-F6 IS-REP. L352 Y ANDEN 3 EST. TORRE PACHECO_v1.0") == "Informe de Seguridad"
    # F6.1 tem de ganhar ao F6, senão o Plan de Pruebas vira Informe de Seguridad.
    assert t("SAG1.AP2-F6.1-PP-REP L352 Y ANDEN 3 EST. TORRE-PACHECO_V0.1") == "Plan de Pruebas / PeS"
    assert t("SAG1-AP2-F7 IPA-REP L352 Y ANDEN 3 EST. TORRE PACHECO_v0.0") == \
        "Informe Previo a la Autorización (F7)"
    assert t("SAG1-AP2-F4.1-FR-RESP001_v0.1") == "Ficha de Requisitos (F4)"
    # Sem o 'de' no meio, como aparece nos envíos reais.
    assert t("INFORME DESVIACIONES SAG1.AP2_v0.3") == "Informe de Desviaciones"


def test_temas_usa_coluna_ou_classifica():
    # Usa a coluna 'tema' quando existe.
    assert guide._temas_da_linha({"tema": "Safety Case, V&V"}) == ["Safety Case", "V&V"]
    # Sem coluna, classifica pelo texto.
    t = guide._temas_da_linha({"hallazgo": "El registro de peligros no tiene medida mitigadora"})
    assert "Hazard Log / REP" in t


def test_contar_ordena_por_criticos():
    rows = [
        {"documento": "Informe de Seguridad", "valoracion": "Crítico", "obra": "A"},
        {"documento": "Informe de Seguridad", "valoracion": "Crítico", "obra": "B"},
        {"documento": "Plan de Seguridad", "valoracion": "Formal", "obra": "A"},
    ]
    for r in rows:
        r["_tipo"] = guide.tipo_documento(r["documento"])
    agg = guide._contar(rows, lambda r: r["_tipo"])
    assert agg[0]["grupo"] == "Informe de Seguridad"   # mais Críticos primeiro
    assert agg[0]["Crítico"] == 2 and agg[0]["Obras"] == 2


def test_exportar_gera_duas_abas():
    import openpyxl

    rows = [
        {"documento": "Caso de Seguridad", "valoracion": "Crítico", "estado": "Abierto",
         "obra": "EXC1", "hallazgo": "falta definición del sistema", "discusion": "",
         "punto": "3.1", "fuente": "LPA-A", "tema": "Safety Case"},
        {"documento": "Apéndice 1. REP", "valoracion": "Importante", "estado": "Cerrado",
         "obra": "EXC2", "hallazgo": "peligro sin mitigación", "discusion": "", "punto": "5",
         "fuente": "LPA-B", "tema": "Hazard Log / REP"},
    ]
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "base.xlsx"
        r = guide.exportar(rows, out)
        assert r["hallazgos"] == 2
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["Por onde começar", "Hallazgos"]
        # a aba Hallazgos tem cabeçalho + 2 linhas
        ws = wb["Hallazgos"]
        assert ws.max_row == 3
        assert ws.cell(row=1, column=1).value == "Tipo doc"
        assert ws.auto_filter.ref is not None


if __name__ == "__main__":
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for fn in fns:
        try:
            fn()
            print(f"PASS {fn.__name__}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"FAIL {fn.__name__}: {e}")
    sys.exit(1 if failed else 0)
